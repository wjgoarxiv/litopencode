#!/usr/bin/env python3
"""Edit existing DOCX files with common operations.

Usage:
    # Find and replace text
    python edit_docx.py input.docx --replace "OLD" "NEW" -o output.docx

    # Append markdown content
    python edit_docx.py input.docx --append-md content.md -o output.docx

    # Insert at placeholder
    python edit_docx.py input.docx --insert-at "{{MARKER}}" "New content" -o output.docx

    # Extract text
    python edit_docx.py input.docx --extract-text

    # Show document structure
    python edit_docx.py input.docx --show-structure

Features:
- Find & Replace across paragraphs, tables, headers, footers
- Append content from Markdown file
- Insert content at placeholder markers
- Extract all text for inspection
- Show document structure (headings)
"""

from __future__ import annotations
import argparse
import re
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.shared import Pt
except ImportError:
    print("Error: 'python-docx' library not found. Install with: pip install python-docx")
    sys.exit(1)


def find_and_replace(doc: Document, old_text: str, new_text: str, case_sensitive: bool = True) -> int:
    """Find and replace text in all document elements. Returns count of replacements."""
    count = 0
    flags = 0 if case_sensitive else re.IGNORECASE
    pattern = re.compile(re.escape(old_text), flags)

    def replace_in_paragraph(paragraph):
        nonlocal count
        full_text = paragraph.text
        if pattern.search(full_text):
            # Count matches
            matches = len(pattern.findall(full_text))
            count += matches

            # Simple approach: reconstruct paragraph
            new_full_text = pattern.sub(new_text, full_text)

            # Clear existing runs and add new text
            # Note: This loses formatting. For format-preserving, need more complex logic.
            for run in paragraph.runs:
                run.text = ""
            if paragraph.runs:
                paragraph.runs[0].text = new_full_text
            else:
                paragraph.add_run(new_full_text)

    # Process main body paragraphs
    for para in doc.paragraphs:
        replace_in_paragraph(para)

    # Process tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    replace_in_paragraph(para)

    # Process headers and footers
    for section in doc.sections:
        for header in [section.header, section.first_page_header, section.even_page_header]:
            if header:
                for para in header.paragraphs:
                    replace_in_paragraph(para)
        for footer in [section.footer, section.first_page_footer, section.even_page_footer]:
            if footer:
                for para in footer.paragraphs:
                    replace_in_paragraph(para)

    return count


def insert_at_marker(doc: Document, marker: str, content: str) -> bool:
    """Insert content at a placeholder marker. Returns True if marker found."""
    found = False

    for para in doc.paragraphs:
        if marker in para.text:
            # Replace marker with content
            for run in para.runs:
                if marker in run.text:
                    run.text = run.text.replace(marker, content)
                    found = True

    # Also check tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    if marker in para.text:
                        for run in para.runs:
                            if marker in run.text:
                                run.text = run.text.replace(marker, content)
                                found = True

    return found


def append_markdown_content(doc: Document, md_path: Path):
    """Append content from a Markdown file to the document."""
    # Import the converter
    script_dir = Path(__file__).parent
    sys.path.insert(0, str(script_dir))

    try:
        from convert_md_to_docx import convert_md_to_docx, html_to_docx
        import markdown
    except ImportError as e:
        print(f"Error importing converter: {e}")
        print("Make sure convert_md_to_docx.py is in the same directory")
        return False

    # Read and convert markdown
    md_text = md_path.read_text(encoding='utf-8')
    md_converter = markdown.Markdown(extensions=[
        'tables',
        'fenced_code',
        'nl2br',
        'sane_lists',
    ])
    html = md_converter.convert(md_text)

    # Add a separator
    doc.add_paragraph()
    doc.add_paragraph('─' * 50)
    doc.add_paragraph()

    # Convert and append HTML content
    html_to_docx(html, doc, md_path)

    return True


def extract_text(doc: Document) -> str:
    """Extract all text from document."""
    lines = []

    for para in doc.paragraphs:
        lines.append(para.text)

    for table in doc.tables:
        lines.append("\n[TABLE]")
        for row in table.rows:
            cells = [cell.text for cell in row.cells]
            lines.append(" | ".join(cells))
        lines.append("[/TABLE]\n")

    return "\n".join(lines)


def show_structure(doc: Document) -> str:
    """Show document structure (headings and paragraphs)."""
    lines = []

    for i, para in enumerate(doc.paragraphs):
        style_name = para.style.name if para.style else "Normal"
        text_preview = para.text[:60] + "..." if len(para.text) > 60 else para.text

        if "Heading" in style_name:
            level = style_name.replace("Heading ", "")
            indent = "  " * (int(level) - 1) if level.isdigit() else ""
            lines.append(f"{indent}[{style_name}] {text_preview}")
        elif para.text.strip():
            lines.append(f"  [{style_name}] {text_preview}")

    # Count tables
    table_count = len(doc.tables)
    if table_count:
        lines.append(f"\n[{table_count} table(s) in document]")

    return "\n".join(lines)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Edit existing DOCX files with common operations"
    )
    parser.add_argument("input", help="Input DOCX file")
    parser.add_argument("-o", "--output", help="Output DOCX file (required for modifications)")

    # Operations (mutually exclusive)
    ops = parser.add_mutually_exclusive_group(required=True)
    ops.add_argument(
        "--replace",
        nargs=2,
        metavar=("OLD", "NEW"),
        help="Find and replace text"
    )
    ops.add_argument(
        "--append-md",
        metavar="MD_FILE",
        help="Append content from Markdown file"
    )
    ops.add_argument(
        "--insert-at",
        nargs=2,
        metavar=("MARKER", "CONTENT"),
        help="Insert content at placeholder marker"
    )
    ops.add_argument(
        "--extract-text",
        action="store_true",
        help="Extract all text from document"
    )
    ops.add_argument(
        "--show-structure",
        action="store_true",
        help="Show document structure (headings)"
    )

    # Options
    parser.add_argument(
        "--case-insensitive",
        action="store_true",
        help="Case-insensitive search for --replace"
    )

    return parser.parse_args()


def main() -> int:
    args = parse_args()

    input_path = Path(args.input).expanduser().resolve()
    if not input_path.exists():
        print(f"Error: Input file not found: {input_path}")
        return 1

    doc = Document(str(input_path))

    # Read-only operations
    if args.extract_text:
        print(extract_text(doc))
        return 0

    if args.show_structure:
        print(show_structure(doc))
        return 0

    # Modification operations require output
    if not args.output:
        print("Error: --output/-o is required for modification operations")
        return 1

    output_path = Path(args.output).expanduser().resolve()

    if args.replace:
        old_text, new_text = args.replace
        count = find_and_replace(doc, old_text, new_text, not args.case_insensitive)
        print(f"Replaced {count} occurrence(s) of '{old_text}' with '{new_text}'")

    elif args.append_md:
        md_path = Path(args.append_md).expanduser().resolve()
        if not md_path.exists():
            print(f"Error: Markdown file not found: {md_path}")
            return 1
        if append_markdown_content(doc, md_path):
            print(f"Appended content from {md_path}")
        else:
            print("Error appending content")
            return 1

    elif args.insert_at:
        marker, content = args.insert_at
        if insert_at_marker(doc, marker, content):
            print(f"Inserted content at marker '{marker}'")
        else:
            print(f"Warning: Marker '{marker}' not found in document")

    # Save modified document
    doc.save(str(output_path))
    print(f"Saved to {output_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
