"""DOCX craft checks with explicit measurement tiers, and the document gate built on them.

`analyse()` is the source-only craft floor (OF-301 prose measure, OF-302 numeric alignment) that
slop_lint's audit route reports. Run as a script, this file is the document gate (`run.py qa`):

  1. Integrity  the package opens, its parts parse and python-docx reopens it
  2. Content    the document has body text, every Markdown heading reached it, no unfilled blank
                ([제품명], XXX, TBD, ○○, {{…}}) is left and no frontmatter leaked into the body
  3. Lint       slop_lint's prose rules on the source (report kind drops the manuscript outline rules),
                its design, Hangul-pairing and craft audit on the DOCX
  4. Output     docx_layout's checks: noun-phrase headings and heading order always, the restraint
                checks (fail under a tonality), with --layout the rendered pages (fill, stranded headings,
                splits, component variety, cover block), with --compare the structural difference
                between two tonalities of one source

    python3 run.py qa draft.docx --source draft.md --layout [--tonality Report] [--kind report]
    python3 run.py qa draft.docx --source draft.md --publisher elsevier --kind manuscript
    python3 run.py qa a.docx --source draft.md --compare b.docx

Prints JSON; exit 0 on pass, 1 on fail. A pass is defect absence, never a verdict on the pages.
"""

from __future__ import annotations

import argparse
import json
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

T = json.loads(Path(__file__).with_name("craft-thresholds.json").read_text())
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
HANGUL = re.compile(r"[가-힣]")


def _text(element):
    return "".join(node.text or "" for node in element.iter(W + "t"))


def _number(text):
    remaining = re.sub(r"[\s,.%+\-±()$€£¥₩▲▼]", "", text.strip())
    return bool(remaining) and remaining.isdigit()


def _attribute(element, name, fallback):
    value = element.get(W + name) if element is not None else None
    return float(value) if value is not None else fallback


def analyse(path):
    with zipfile.ZipFile(path) as archive:
        document = ET.fromstring(archive.read("word/document.xml"))
        styles = ET.fromstring(archive.read("word/styles.xml"))
    body = document.find(W + "body")
    section = body.find(W + "sectPr") if body is not None else None
    page = section.find(W + "pgSz") if section is not None else None
    margins = section.find(W + "pgMar") if section is not None else None
    width = _attribute(page, "w", T["fallbackPageWidthTwips"])
    left = _attribute(margins, "left", T["fallbackMarginTwips"])
    right = _attribute(margins, "right", T["fallbackMarginTwips"])
    usable_inches = max(0, width - left - right) / 1440
    normal = next((style for style in styles.findall(W + "style") if style.get(W + "styleId") == "Normal"), None)
    normal_size = normal.find(".//" + W + "sz") if normal is not None else None
    font_pt = _attribute(normal_size, "val", T["fallbackFontPt"] * 2) / 2
    findings = []

    if body is None:
        return findings
    for paragraph_no, paragraph in enumerate(body.findall(W + "p"), 1):
        text = _text(paragraph).strip()
        if len(text) < 40:
            continue
        letters = [char for char in text if not char.isspace()]
        cjk = bool(letters) and len(HANGUL.findall(text)) / len(letters) >= T["cjkMajority"]
        factor = T["cjkGlyphFactor"] if cjk else T["latinGlyphFactor"]
        ceiling = T["cjkMeasureMax"] if cjk else T["latinMeasureMax"]
        run_sizes = [float(node.get(W + "val")) / 2 for node in paragraph.findall(".//" + W + "sz") if node.get(W + "val")]
        size = max(run_sizes, default=font_pt)
        measure = usable_inches * 72 / max(1, size * factor)
        if measure > ceiling:
            findings.append({"rule": "OF-301", "severity": "MEDIUM", "tier": "derived", "where": f"paragraph {paragraph_no}", "value": round(measure, 1), "threshold": ceiling})

    for table_no, table in enumerate(body.iter(W + "tbl"), 1):
        rows = table.findall(W + "tr")
        if len(rows) < 3:
            continue
        data = [[cell for cell in row.findall(W + "tc")] for row in rows[1:]]
        columns = max(map(len, data), default=0)
        for column in range(columns):
            cells = [row[column] for row in data if len(row) > column]
            if len(cells) < 2 or sum(_number(_text(cell)) for cell in cells) / len(cells) < T["numericMajority"]:
                continue
            for row_no, cell in enumerate(cells, 2):
                if not _number(_text(cell)):
                    continue
                for paragraph in cell.findall(W + "p"):
                    if not _text(paragraph).strip():
                        continue
                    alignment = paragraph.find("./" + W + "pPr/" + W + "jc")
                    value = alignment.get(W + "val") if alignment is not None else None
                    if value != "right":
                        findings.append({"rule": "OF-302", "severity": "HIGH" if value else "MEDIUM", "tier": "measured" if value else "derived", "where": f"table {table_no} row {row_no} col {column + 1}", "value": value, "threshold": "right"})
    return findings


HEADING = re.compile(r"^(#{1,6})\s+(.+?)\s*#*\s*$")
# A bracket holding a digit is a citation ([1], [S3]); a blank holds none.
BLANK = re.compile(r"\[(?![^\]\n]*\d)[^\]\n]{1,24}\]|\bX{3,}\b|\bTBD\b|\bTODO\b|○○|OOO|\{\{[^}]*\}\}|_{4,}|＿{2,}")
LEAKED_FRONTMATTER = re.compile(r"^(?:---|title:|author:|authors:|date:|subtitle:)\s*", re.M)
# Under the plain profile only the rules every profile shares apply to the Word file.
SHARED_DESIGN_RULES = {"rule-45-highlighted-text", "rule-46-tracked-changes", "rule-55-ellipsis", "rule-62-numeric-alignment"}


def _integrity(path):
    problems = []
    try:
        with zipfile.ZipFile(path) as archive:
            names = set(archive.namelist())
            for part in ("[Content_Types].xml", "word/document.xml", "word/styles.xml"):
                if part not in names:
                    problems.append(f"missing {part}")
            for name in names:
                if name.endswith((".xml", ".rels")):
                    try:
                        ET.fromstring(archive.read(name))
                    except ET.ParseError as error:
                        problems.append(f"{name}: {error}")
    except (OSError, zipfile.BadZipFile) as error:
        problems.append(str(error))
    if not problems:
        from docx import Document
        try:
            Document(str(path))
        except Exception as error:  # python-docx raises many kinds on a broken part
            problems.append(f"python-docx cannot reopen it: {error}")
    return {"pass": not problems, "problems": problems[:20]}


def _headings(markdown_text):
    found, fenced = [], False
    for line in markdown_text.splitlines():
        if line.lstrip().startswith(("```", "~~~")):
            fenced = not fenced
            continue
        match = None if fenced else HEADING.match(line)
        if match:
            found.append(re.sub(r"[*_`]", "", match.group(2)).strip())
    return found


def _content(path, markdown_text, layout):
    from docx import Document
    document = Document(str(path))
    paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    cells = [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    blanks = sorted({match.group(0) for text in paragraphs + cells for match in BLANK.finditer(text)})
    leaked = [text for text in paragraphs[:8] if LEAKED_FRONTMATTER.match(text)]
    missing = []
    if markdown_text:
        # A heading may stand inside a side-by-side columns row, so table text counts too.
        text = layout.norm("".join(paragraphs + cells))
        missing = [heading for heading in _headings(markdown_text) if layout.norm(heading) not in text]
    chars = sum(map(len, paragraphs)) + sum(map(len, cells))
    return {"pass": chars > 0 and not missing and not blanks and not leaked, "blanks": blanks[:20],
            "frontmatter_leak": leaked[:4], "missing_headings": missing[:20], "paragraphs": len(paragraphs),
            "tables": len(document.tables), "body_chars": chars}


def _lint(source, path, publisher_name, kind, explicit):
    import slop_lint
    here = Path(__file__).resolve().parent
    publisher = slop_lint.load_publisher(here.parent / "templates" / "registry.yaml", publisher_name)
    prose, locale = [], "auto"
    if source is not None:
        phrases = slop_lint.load_phrase_rules(here.parent / "references" / "slop_phrase_list.yaml")
        prose, _front, _body, locale = slop_lint.lint_text(source.read_text(encoding="utf-8"), publisher, phrases, "auto")
        if kind == "report":
            prose = [item for item in prose if item.rule_id not in slop_lint.MANUSCRIPT_ONLY_RULES]
        elif not explicit:
            prose = [item for item in prose if item.rule_id != "rule-08-structure-order"]
    design = slop_lint.audit_docx_design(path, publisher)
    if locale in {"ko", "mixed"} or source is None:
        design += slop_lint.audit_docx_cjk(path)
    craft = slop_lint.audit_docx_craft(path)
    design += [item for item in craft if item.severity == "HIGH"]
    if not explicit:
        design = [item for item in design if item.rule_id in SHARED_DESIGN_RULES]
    entry = lambda item: {"rule": item.rule_id, "line": item.line, "section": item.section, "message": item.message, "excerpt": item.excerpt}  # noqa: E731
    return {"pass": not prose and not design, "locale": locale, "prose": [entry(item) for item in prose][:40],
            "design": [entry(item) for item in design][:40], "advisories": [entry(item) for item in craft if item.severity != "HIGH"][:40]}


def _render(path, folder):
    """The document as PDF through soffice with a private profile, or (None, reason)."""
    office = shutil.which("soffice") or shutil.which("libreoffice")
    if office is None:
        return None, "soffice is not on PATH; the page checks did not run"
    profile = Path(folder) / "profile"
    command = [office, f"-env:UserInstallation=file://{profile}", "--headless", "--convert-to", "pdf", "--outdir", str(folder), str(path)]
    try:
        subprocess.run(command, capture_output=True, timeout=180, check=False)
    except subprocess.TimeoutExpired:
        return None, "soffice did not finish within 180 s; the page checks did not run"
    pdf = Path(folder) / (Path(path).stem + ".pdf")
    return (pdf, "soffice") if pdf.exists() else (None, "soffice wrote no PDF; the page checks did not run")


def _output(path, source, args, layout):
    from convert_md_to_docx import parse_frontmatter
    front = parse_frontmatter(source.read_text(encoding="utf-8"))[0] if source else {}
    tonality = bool(args.tonality or front.get("tonality"))
    info = layout.read_docx(path)
    found = layout.structural(info, front, args.kind == "manuscript", tonality) + layout.restraint(path, tonality)
    result = {"components": info["components"], "rendered": False}
    if args.compare:
        mine, other = layout.structure(path), layout.structure(args.compare.expanduser().resolve())
        differ = layout.structure_diff(mine, other)
        result["structure"] = {"this": mine, "other": other, "differ": differ}
        if len(differ) < 3:
            found.append(layout.finding("tonality.structure", "FAIL", None,
                                        f"the two documents differ in {len(differ)} structural feature(s): {', '.join(differ) or 'none'} (at least 3)",
                                        "Set the tonalities apart in structure: title block, summary form, numbering, component set, running head, contents."))
    if args.layout:
        with tempfile.TemporaryDirectory(prefix="litopencode-docx-pages-") as folder:
            pdf, how = _render(path, folder)
            if pdf is None:
                result["render"] = how
            else:
                page_found, pages = layout.paged(info, layout.read_pdf(pdf), tonality, front)
                found += page_found
                result.update(pages, rendered=True)
    result["findings"] = [item for item in found if item["severity"] == "FAIL"]
    result["advisories"] = [item for item in found if item["severity"] != "FAIL"]
    return result


def main(argv=None):
    parser = argparse.ArgumentParser(description="Pass/fail gate for a converted DOCX.")
    parser.add_argument("docx", type=Path)
    parser.add_argument("--source", type=Path, default=None, help="the Markdown the DOCX was converted from")
    parser.add_argument("--publisher", default=None, help="the publisher profile it was converted with; omit for a tonality or the plain profile")
    parser.add_argument("--kind", choices=("report", "manuscript"), default="report")
    parser.add_argument("--layout", action="store_true", help="render through soffice and check the pages")
    parser.add_argument("--tonality", default=None, help="the tonality it was converted with, when the frontmatter does not name it")
    parser.add_argument("--compare", type=Path, default=None, help="the same source in another tonality; fail under three structural differences")
    args = parser.parse_args(argv)
    sys.dont_write_bytecode = True
    import docx_layout as layout

    path = args.docx.expanduser().resolve()
    source = args.source.expanduser().resolve() if args.source else None
    integrity = _integrity(path)
    if integrity["pass"]:
        content = _content(path, source.read_text(encoding="utf-8") if source else None, layout)
        lint = _lint(source, path, args.publisher or "korean-generic", args.kind, args.publisher is not None)
        output = _output(path, source, args, layout)
    else:
        content, lint, output = {"pass": False}, {"pass": False}, {"findings": [], "advisories": []}
    reasons = []
    if not integrity["pass"]:
        reasons.append("integrity: " + "; ".join(integrity["problems"][:5]))
    if content.get("blanks"):
        reasons.append(f"content: unfilled blanks {content['blanks'][:5]} — fill them with realistic values labelled as examples or assumptions")
    if content.get("frontmatter_leak"):
        reasons.append("content: frontmatter printed as body text — convert with run.py md-to-docx, which reads it")
    if integrity["pass"] and not content["pass"] and not content["blanks"] and not content["frontmatter_leak"]:
        reasons.append(f"content: empty document or headings missing from the output {content['missing_headings'][:5]}")
    if lint.get("prose"):
        reasons.append(f"prose lint: {len(lint['prose'])} finding(s)")
    if lint.get("design"):
        reasons.append(f"design audit: {len(lint['design'])} finding(s)")
    if output["findings"]:
        reasons.append("output checks: " + ", ".join(sorted({item["check"] for item in output["findings"]})))
    report = {"file": str(path), "pass": not reasons, "failure_reasons": reasons, "profile": args.publisher or "plain",
              "kind": args.kind, "integrity": integrity, "content": content, "lint": lint, "output": output}
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0 if report["pass"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
