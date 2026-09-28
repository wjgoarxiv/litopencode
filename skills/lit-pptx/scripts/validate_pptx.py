#!/usr/bin/env python3
"""
validate_pptx.py — lightweight anti-slop validator for PPTX outputs.

Checks deck text hygiene, title presence, placeholder/default copy, forbidden
terms, and coarse visual consistency signals. Prints JSON.
"""

from __future__ import annotations

import json
import re
import sys
from collections import Counter
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER


SKILL_DIR = Path(__file__).resolve().parent.parent  # repo root (validate_pptx.py lives in scripts/)
FORBIDDEN_PATH = SKILL_DIR / "FORBIDDEN_TERMS.json"
DEFAULT_BAD_FONTS = {"Calibri", "Aptos"}
DEFAULT_PLACEHOLDER_TEXT = {
    "click to add title",
    "click to add subtitle",
    "click to add text",
    "title",
    "subtitle",
}
MAX_DISTINCT_FILLS = 8
WEIGHTS = {
    "missing_titles": 20,
    "forbidden_terms": 20,
    "default_placeholder_text": 20,
    "default_fonts": 15,
    "rainbow_fills": 10,
}


def _load_terms() -> tuple[list[str], list[str], bool, bool]:
    if not FORBIDDEN_PATH.exists():
        return [], [], True, False
    data = json.loads(FORBIDDEN_PATH.read_text())
    return (
        data.get("terms", []),
        data.get("soft_terms", []),
        bool(data.get("case_insensitive", True)),
        bool(data.get("whole_word", False)),
    )


def _build_regex(terms: list[str], case_insensitive: bool, whole_word: bool) -> re.Pattern[str] | None:
    if not terms:
        return None
    escaped = [re.escape(term) for term in terms]
    pattern = "|".join(escaped)
    if whole_word:
        pattern = rf"(?:(?<=\W)|^)(?:{pattern})(?:(?=\W)|$)"
    flags = re.IGNORECASE if case_insensitive else 0
    return re.compile(pattern, flags)


def _shape_text(shape) -> str:
    if not getattr(shape, "has_text_frame", False):
        return ""
    parts: list[str] = []
    for paragraph in shape.text_frame.paragraphs:
        chunk = "".join(run.text for run in paragraph.runs).strip()
        if chunk:
            parts.append(chunk)
    return "\n".join(parts).strip()


def _shape_fill(shape) -> str | None:
    fill = getattr(shape, "fill", None)
    if fill is None or fill.type is None:
        return None
    try:
        fore = getattr(fill, "fore_color", None)
    except TypeError:
        return None
    rgb = getattr(fore, "rgb", None)
    if isinstance(rgb, RGBColor):
        return str(rgb)
    if isinstance(rgb, str):
        return rgb.upper()
    return None


def _iter_runs(slide):
    for shape in slide.shapes:
        if not getattr(shape, "has_text_frame", False):
            continue
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                yield shape, run


def _slide_has_title(slide, slide_height: int | None) -> bool:
    title_shape = getattr(slide.shapes, "title", None)
    if title_shape is not None and _shape_text(title_shape):
        return True
    for shape in slide.shapes:
        if not getattr(shape, "is_placeholder", False):
            continue
        placeholder = getattr(shape, "placeholder_format", None)
        if placeholder and placeholder.type == PP_PLACEHOLDER.TITLE and _shape_text(shape):
            return True
    text_shapes = []
    for shape in slide.shapes:
        if not getattr(shape, "has_text_frame", False):
            continue
        text = _shape_text(shape)
        if not text:
            continue
        text_shapes.append(shape)
        top = getattr(shape, "top", None)
        if slide_height is not None and top is not None and top <= slide_height * 0.22:
            return True
    if len(text_shapes) <= 2 and any(_shape_text(shape) for shape in text_shapes):
        return True
    return False


def lint(path: str | Path) -> dict[str, Any]:
    prs = Presentation(str(path))
    hard_terms, soft_terms, case_insensitive, whole_word = _load_terms()
    hard_re = _build_regex(hard_terms, case_insensitive, whole_word)
    soft_re = _build_regex(soft_terms, case_insensitive, whole_word)

    missing_titles: list[int] = []
    placeholder_hits: list[str] = []
    hard_hits: list[str] = []
    soft_hits: list[str] = []
    default_font_hits: list[str] = []
    fills: Counter[str] = Counter()

    for slide_idx, slide in enumerate(prs.slides, start=1):
        if not _slide_has_title(slide, prs.slide_height):
            missing_titles.append(slide_idx)

        for shape_idx, shape in enumerate(slide.shapes, start=1):
            fill = _shape_fill(shape)
            if fill:
                fills[fill] += 1

            text = _shape_text(shape)
            if not text:
                continue

            lowered = text.strip().lower()
            if lowered in DEFAULT_PLACEHOLDER_TEXT:
                placeholder_hits.append(f"slide {slide_idx} shape {shape_idx}: {text[:80]}")

            if hard_re and hard_re.search(text):
                hard_hits.append(f"slide {slide_idx} shape {shape_idx}: {text[:120]}")
            elif soft_re and soft_re.search(text):
                soft_hits.append(f"slide {slide_idx} shape {shape_idx}: {text[:120]}")

        for shape, run in _iter_runs(slide):
            font_name = (run.font.name or "").strip()
            if font_name in DEFAULT_BAD_FONTS:
                default_font_hits.append(f"slide {slide_idx}: {font_name}")

    last_slide_text = "\n".join(_shape_text(shape) for shape in prs.slides[-1].shapes if getattr(shape, "has_text_frame", False)).lower() if prs.slides else ""
    has_source_note = "source:" in last_slide_text or "출처" in last_slide_text

    failures = {
        "missing_titles": bool(missing_titles),
        "forbidden_terms": bool(hard_hits),
        "default_placeholder_text": bool(placeholder_hits),
        "default_fonts": bool(default_font_hits),
        "rainbow_fills": len(fills) > MAX_DISTINCT_FILLS,
    }
    slop_score = sum(WEIGHTS[key] for key, failed in failures.items() if failed)

    return {
        "file": str(path),
        "slide_count": len(prs.slides),
        "pass": slop_score <= 10 and not hard_hits and not placeholder_hits,
        "slop_score": slop_score,
        "checks": {
            "missing_titles": {"pass": not missing_titles, "slides": missing_titles},
            "forbidden_terms": {"pass": not hard_hits, "hard_hits": hard_hits[:20], "soft_warnings": soft_hits[:20]},
            "default_placeholder_text": {"pass": not placeholder_hits, "hits": placeholder_hits[:20]},
            "default_fonts": {"pass": not default_font_hits, "hits": default_font_hits[:20]},
            "rainbow_fills": {"pass": len(fills) <= MAX_DISTINCT_FILLS, "distinct_fills": len(fills), "sample": list(fills.keys())[:12]},
            "missing_source_note": {"pass": has_source_note},
        },
    }


def main(argv: list[str]) -> int:
    if len(argv) != 2:
        print(f"Usage: {Path(argv[0]).name} <file.pptx>", file=sys.stderr)
        return 2
    report = lint(argv[1])
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0 if report["pass"] else 1


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
