"""Measured OOXML craft checks for the LitOpenCode slide QA gate.

  OF-101  at most four accent hue families per slide (among solid shape fills); five fail
  OF-120  a hard offset shadow (no blur) only in a pack that allows one, and there every shadow is hard

A tonality pack may widen two limits through its own `qa:` block: `accent-families` and
`display-accent-families` (how many accent hue families a content or display slide may carry before
OF-101 fails; five or six, the default is four) and `offset-shadow` (OF-120). The gate finds the pack
through the deck's tonality stamp; a deck without a stamp keeps the defaults.
"""

from __future__ import annotations

import colorsys
import json
import re
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN

from deck_output import check_figures, check_labels, check_output, check_run_ins

T = json.loads(Path(__file__).with_name("craft-thresholds.json").read_text())
EMU = 914400
NUMERIC = re.compile(r"^[\d\s,.%+\-±()$€£¥₩▲▼]+$")
HANGUL = re.compile(r"[가-힣]")
EMOJI = re.compile(r"[\U0001F000-\U0001FAFF\u2600-\u27BF]")
STAMP = re.compile(r"\blit-pptx tonality=(\S+) density=(\d+) grid=(\w+) variance=(\d+)")
QA_KEY = re.compile(r"^  ([a-z-]+):\s*(\S+)\s*$")
# Families the tonality compiler draws as display pages; every other `family@` name is a content slide.
DISPLAY_FAMILY = re.compile(r"(?:cover|section)(?:-|$)|(?:statement|quote|closing-statement|image-full)$")


def _box(shape):
    return tuple(getattr(shape, key) / EMU for key in ("left", "top", "width", "height"))


def _gap(a, b):
    return max(0, max(a[0], b[0]) - min(a[0] + a[2], b[0] + b[2]),
               max(a[1], b[1]) - min(a[1] + a[3], b[1] + b[3]))


def _contains(outer, inner):
    return outer[0] <= inner[0] and outer[1] <= inner[1] and outer[0] + outer[2] >= inner[0] + inner[2] and outer[1] + outer[3] >= inner[1] + inner[3]


def _fill(shape):
    try:
        return tuple(shape.fill.fore_color.rgb) if shape.fill.type is not None else None
    except (AttributeError, TypeError):
        return None


def _hue(rgb):
    if rgb is None:
        return None
    hue, saturation, lightness = colorsys.rgb_to_hls(*(channel / 255 for channel in rgb))
    if saturation < T["accentSaturationMin"] or not T["accentLightnessMin"] <= lightness <= T["accentLightnessMax"]:
        return None
    return hue * 360


def _family_count(hues):
    families = []
    for hue in hues:
        if not any(min(abs(hue - other), 360 - abs(hue - other)) <= T["accentHueBucketDegrees"] for other in families):
            families.append(hue)
    return len(families)


def _text_size(shape):
    if not getattr(shape, "has_text_frame", False):
        return 0
    return max((run.font.size.pt for paragraph in shape.text_frame.paragraphs for run in paragraph.runs if run.font.size), default=0)


def _display(slide):
    family = next((shape.name.split("@", 1)[1] for shape in slide.shapes if shape.name.startswith("family@")), None)
    if family is not None:
        return bool(DISPLAY_FAMILY.match(family))
    text = [shape for shape in slide.shapes if getattr(shape, "has_text_frame", False) and shape.text.strip()]
    visual = [shape for shape in slide.shapes if getattr(shape, "has_table", False) or getattr(shape, "has_chart", False)]
    return bool(text) and max(map(_text_size, text)) >= T["displayPt"] and len(text) <= 2 and not visual


def _numeric(text):
    stripped = re.sub(r"[\s,.%+\-±()$€£¥₩▲▼]", "", text.strip())
    return bool(stripped) and stripped.isdigit()


def _radius(shape):
    try:
        return float(shape.adjustments[0]) * min(shape.width, shape.height) / EMU
    except (AttributeError, IndexError, TypeError, ValueError):
        return None


def _page_furniture(box, width, height, tolerance=0.03):
    """A fill on two or more canvas edges (a band, a rail, a colour field) frames the page, not content."""
    x, y, w, h = box
    return sum([x <= tolerance, y <= tolerance, x + w >= width - tolerance, y + h >= height - tolerance]) >= 2


def _vertical_gaps(top, bottom, boxes):
    cursor = top
    gaps = []
    for _, y, _, height in sorted(boxes, key=lambda box: box[1]):
        gaps.append(max(0, y - cursor))
        cursor = max(cursor, y + height)
    gaps.append(max(0, bottom - cursor))
    return gaps


def pack_allowances(deck):
    """The stamped tonality pack's `qa:` block as {key: int or bool}; empty without a stamp or a block."""
    stamp = STAMP.search(deck.core_properties.subject or "")
    if not stamp or not re.fullmatch(r"[a-z][a-z0-9-]*", stamp.group(1)):
        return {}
    pack = Path(__file__).resolve().parent.parent / "templates" / "tonalities" / stamp.group(1) / "pack.yaml"
    if not pack.is_file():
        return {}
    allowed, inside = {}, False
    for line in pack.read_text(encoding="utf-8").splitlines():
        if not line.strip() or line.lstrip().startswith("#"):
            continue
        if not line.startswith(" "):
            inside = line.rstrip() == "qa:"
            continue
        match = QA_KEY.match(line) if inside else None
        if match:
            value = match.group(2)
            allowed[match.group(1)] = value == "true" if value in ("true", "false") else int(value) if value.isdigit() else value
    return allowed


def _accent_limits(display, allowed):
    """The hue-family count that is advisory and the count that fails. A pack's own limit only widens the
    default: N families pass, N + 1 fail, with no advisory step between."""
    medium = T["displayAccentMedium"] if display else T["contentAccentMedium"]
    high = T["displayAccentHigh"] if display else T["contentAccentHigh"]
    allowance = allowed.get("display-accent-families" if display else "accent-families")
    if isinstance(allowance, int) and not isinstance(allowance, bool):
        medium = high = max(high, allowance + 1)
    return medium, high


def analyse(path):
    deck = Presentation(str(path))
    allowed = pack_allowances(deck)
    canvas = deck.slide_width * deck.slide_height
    findings = check_output(deck) + check_labels(deck) + check_figures(deck) + check_run_ins(deck)

    def add(rule, severity, slide, value, threshold, tier="measured", shape=None):
        findings.append({"rule": rule, "severity": severity, "tier": tier, "slide": slide,
                         "shape": shape, "value": value, "threshold": threshold})

    for slide_no, slide in enumerate(deck.slides, 1):
        # The sample-data tag is a deliberate label on every slide, not part of the design.
        shapes = [shape for shape in slide.shapes if isinstance(shape.width, int) and isinstance(shape.height, int)
                  and not shape.name.startswith("lit-notice")]
        display = _display(slide)
        hues = []
        for index, shape in enumerate(shapes, 1):
            if shape.width * shape.height >= T["fullBleedFraction"] * canvas:
                continue
            hue = _hue(_fill(shape))
            if hue is not None:
                hues.append(hue)
            if shape._element.xpath('.//a:glow'):
                add("OF-107", "HIGH", slide_no, "glow", "none", shape=index)
            if shape._element.xpath('.//a:rPr/a:gradFill'):
                add("OF-106", "HIGH", slide_no, "gradient text", "none", shape=index)
            for bullet in shape._element.xpath('.//a:pPr/a:buChar'):
                char = bullet.get("char", "")
                if EMOJI.search(char):
                    add("OF-108", "HIGH", slide_no, char, "no emoji bullet", shape=index)

            if getattr(shape, "has_text_frame", False) and shape.text.strip():
                size = _text_size(shape)
                if not T["captionPt"] < size < T["displayPt"]:
                    continue
                chars = len(re.sub(r"\s", "", shape.text))
                if not chars:
                    continue
                usable = max(1, shape.width - shape.text_frame.margin_left - shape.text_frame.margin_right)
                estimated_width = sum((0.98 if HANGUL.match(char) else 0.56) * size for char in shape.text) / 72
                lines = max(shape.text.count("\n") + 1, round(estimated_width / max(usable / EMU, .01) + .5))
                if lines < 2:
                    continue
                measure = chars / lines
                cjk = len(HANGUL.findall(shape.text)) / chars >= T["cjkMajority"]
                ceiling = T["bodyCjkCharsMax"] if cjk else T["bodyLatinCharsMax"]
                if measure > ceiling:
                    add("OF-102", "HIGH", slide_no, round(measure, 1), ceiling, "derived", index)
                elif lines >= T["bodyNarrowLines"] and measure < T["bodyNarrowChars"]:
                    add("OF-102", "MEDIUM", slide_no, round(measure, 1), T["bodyNarrowChars"], "derived", index)

            if getattr(shape, "has_table", False):
                table = shape.table
                for col in range(len(table.columns)):
                    cells = [table.cell(row, col) for row in range(1, len(table.rows))]
                    if len(cells) < 2 or sum(_numeric(cell.text) for cell in cells) / len(cells) < T["numericMajority"]:
                        continue
                    for row, cell in enumerate(cells, 2):
                        if not _numeric(cell.text):
                            continue
                        for paragraph in cell.text_frame.paragraphs:
                            if not paragraph.text.strip():
                                continue
                            if paragraph.alignment is None:
                                add("OF-103", "MEDIUM", slide_no, {"row": row, "col": col + 1, "alignment": None}, "right", "derived", index)
                            elif paragraph.alignment != PP_ALIGN.RIGHT:
                                add("OF-103", "HIGH", slide_no, {"row": row, "col": col + 1, "alignment": str(paragraph.alignment)}, "right", shape=index)

        count = _family_count(hues)
        medium, high = _accent_limits(display, allowed)
        if count >= medium:
            add("OF-101", "HIGH" if count >= high else "MEDIUM", slide_no, count, {"medium": medium, "high": high})

        # OF-120: a hard offset shadow (blur 0) is a pack's look or nothing; a pack that has it keeps every shadow hard.
        for index, shape in enumerate(slide.shapes, 1):
            for shadow in shape._element.xpath('.//a:outerShdw'):
                hard = int(shadow.get("blurRad") or 0) == 0
                if hard and not allowed.get("offset-shadow"):
                    add("OF-120", "HIGH", slide_no, "hard offset shadow in a pack that does not draw one", "no shadow", shape=index)
                elif not hard and allowed.get("offset-shadow"):
                    add("OF-120", "HIGH", slide_no, "soft shadow beside the pack's hard offset shadows", "blur 0", shape=index)

        rounds = [(index, shape, _box(shape), _radius(shape)) for index, shape in enumerate(shapes, 1) if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and "ROUNDED_RECTANGLE" in str(shape.auto_shape_type)]
        for inner_index, inner, ibox, iradius in rounds:
            for outer_index, outer, obox, oradius in rounds:
                if inner is outer or iradius is None or oradius is None or not _contains(obox, ibox):
                    continue
                insets = [ibox[0] - obox[0], obox[0] + obox[2] - ibox[0] - ibox[2], ibox[1] - obox[1], obox[1] + obox[3] - ibox[1] - ibox[3]]
                if max(insets) - min(insets) > T["radiusInsetToleranceIn"] or oradius <= T["outerRadiusMinIn"]:
                    continue
                expected = max(0, oradius - sum(insets) / len(insets))
                if abs(iradius - expected) > T["radiusDifferenceToleranceIn"]:
                    add("OF-104", "HIGH", slide_no, round(iradius, 3), round(expected, 3), shape=inner_index)

        panels = [(shape, _box(shape)) for shape in shapes if _fill(shape) is not None and .2 < shape.width * shape.height / canvas < T["fullBleedFraction"]]
        groups = []
        for panel, pbox in panels:
            members = [_box(item) for item in shapes if item is not panel and getattr(item, "has_text_frame", False) and item.text.strip() and _contains(pbox, _box(item))]
            if len(members) >= 2:
                groups.append((pbox, members))
        if len(groups) >= 2:
            intra = max((min(_gap(a, b) for i, a in enumerate(members) for b in members[i + 1:]) for _, members in groups), default=0)
            inter = min((_gap(a, b) for i, (a, _) in enumerate(groups) for b, _ in groups[i + 1:]), default=0)
            if intra > 0 and inter < T["groupRatioMin"] * intra:
                add("OF-105", "MEDIUM", slide_no, round(inter / intra, 2), T["groupRatioMin"])

        if not display:
            # The compiler names its title frame; a title in the lower half closes the body from below.
            title = next((shape for shape in shapes if shape.name.startswith("title@")), None) or max(
                (shape for shape in shapes if getattr(shape, "has_text_frame", False) and shape.text.strip() and _text_size(shape) >= T["displayPt"]), key=_text_size, default=None)
            if title is not None and title.top > deck.slide_height / 2:
                top, bottom = .8, title.top / EMU - .1
            else:
                top = (title.top + title.height) / EMU + .1 if title else .8
                bottom = deck.slide_height / EMU - .55
            held = [_box(shape) for shape in shapes if shape is not title and (getattr(shape, "has_table", False) or getattr(shape, "has_chart", False) or shape.shape_type == MSO_SHAPE_TYPE.PICTURE or (getattr(shape, "has_text_frame", False) and shape.text.strip() and _text_size(shape) > T["captionPt"]))]
            # A title in a side rail: the body stands beside the rail from the title's top, and what stands
            # in the rail under the title (criteria labels, takeaways, the source) belongs to the rail.
            if title is not None and title.left < .1 * deck.slide_width and title.width < .35 * deck.slide_width and title.top < deck.slide_height / 4:
                rx, ry, rw, rh = _box(title)
                beside = [box for box in held if box[0] >= rx + rw - .05 and box[1] < bottom]
                under = [box for box in held if box[0] + box[2] <= rx + rw + .05 and box[1] >= ry + rh - .05]
                if beside and len(beside) + len(under) == len([box for box in held if box[1] < bottom]):
                    top, held = ry, beside
            content = [box for box in held if box[1] < bottom and box[1] + box[3] > top]
            if content and bottom > top:
                gaps = _vertical_gaps(top, bottom, content)
                fraction = max(gaps) / (bottom - top)
                if fraction >= T["emptyBandMedium"]:
                    add("OF-109", "HIGH" if fraction > T["emptyBandHigh"] else "MEDIUM", slide_no, round(fraction, 3), {"medium": T["emptyBandMedium"], "high": T["emptyBandHigh"]})

        card_gaps = []
        for index, card in enumerate([] if display else shapes, 1):
            if _fill(card) is None or card.width * card.height / canvas >= T["fullBleedFraction"]:
                continue
            x, y, width, height = _box(card)
            if width <= 0 or height <= 0 or _page_furniture((x, y, width, height), deck.slide_width / EMU, deck.slide_height / EMU):
                continue
            frame = getattr(card, "text_frame", None) if getattr(card, "has_text_frame", False) else None
            top = y + (frame.margin_top / EMU if frame else 0)
            bottom = y + height - (frame.margin_bottom / EMU if frame else 0)
            inner = [(item, _box(item)) for item in shapes if item is not card and _contains((x, y, width, height), _box(item)) and (getattr(item, "has_table", False) or getattr(item, "has_chart", False) or item.shape_type == MSO_SHAPE_TYPE.PICTURE or getattr(item, "has_text_frame", False) and item.text.strip())]
            if not inner or bottom <= top:
                continue
            text_items = [item for item, _ in inner if getattr(item, "has_text_frame", False) and item.text.strip()]
            visuals = [item for item, _ in inner if getattr(item, "has_table", False) or getattr(item, "has_chart", False) or item.shape_type == MSO_SHAPE_TYPE.PICTURE]
            if text_items and len(text_items) <= 2 and not visuals and max(map(_text_size, text_items)) >= T["displayPt"]:
                continue
            gaps = _vertical_gaps(top, bottom, [box for _, box in inner])
            card_gaps.append({"index": index, "top": y, "height": height, "head": max(gaps[:-1], default=0) / (bottom - top), "tail": gaps[-1] / (bottom - top)})
        for card in card_gaps:
            row = [peer for peer in card_gaps if abs(peer["top"] - card["top"]) <= T["equalRowToleranceIn"] and abs(peer["height"] - card["height"]) <= T["equalRowToleranceIn"]]
            trailing = min(peer["tail"] for peer in row)
            fraction = max(card["head"], trailing)
            if fraction >= T["emptyBandMedium"]:
                add("OF-109", "HIGH" if fraction > T["emptyBandHigh"] else "MEDIUM", slide_no, round(fraction, 3), {"medium": T["emptyBandMedium"], "high": T["emptyBandHigh"]}, shape=card["index"])

    return {"pass": not any(item["severity"] == "HIGH" for item in findings), "findings": findings}
