#!/usr/bin/env python3
"""Inspect authored slide objects for off-canvas and text-frame overflow.

The report shape is consumed by qa_deck.py. Decorative backgrounds and intentional
layering are not counted as defects; semantic overlap is checked by that gate.
"""

import argparse
import json
from functools import lru_cache
from pathlib import Path

from PIL import ImageFont
from pptx import Presentation
from pptx.util import Inches

FACES = Path(__file__).resolve().parents[1] / "pretendard-font" / "public" / "static"


@lru_cache(maxsize=8)
def _face(bold):
    return ImageFont.truetype(str(FACES / f"Pretendard-{'Bold' if bold else 'Regular'}.otf"), 100)


def _measured_lines(line, size, bold, usable_points):
    """Lines a run of the bundled face takes at this width, wrapped between words like the renderer."""
    face = _face(bold)
    width = lambda text: face.getlength(text) * size / 100
    lines, current = 1, ""
    for word in line.split(" "):
        candidate = f"{current} {word}" if current else word
        if current and width(candidate) > usable_points:
            lines += 1
            current = word
        else:
            current = candidate
        while width(current) > usable_points and len(current) > 1:
            lines += 1
            current = current[max(1, int(len(current) * usable_points / width(current))):]
    return lines


def inspect(path):
    deck = Presentation(path)
    report = {}
    tolerance = Inches(0.04)
    for slide_number, slide in enumerate(deck.slides, 1):
        issues = {}
        for index, shape in enumerate(slide.shapes, 1):
            if not (shape.has_text_frame or shape.shape_type == 13):
                continue
            name = f"shape_{index}"
            overflow = {}
            off_slide = (shape.left < -tolerance or shape.top < -tolerance or
                         shape.left + shape.width > deck.slide_width + tolerance or
                         shape.top + shape.height > deck.slide_height + tolerance)
            if off_slide:
                visible_width = max(0, min(shape.left + shape.width, deck.slide_width) - max(shape.left, 0))
                visible_height = max(0, min(shape.top + shape.height, deck.slide_height) - max(shape.top, 0))
                visible_fraction = visible_width * visible_height / max(1, shape.width * shape.height)
                # The enrolled cover deliberately bleeds decorative pictures beyond
                # the canvas. Record a partial bleed, but fail a nearly lost picture.
                if shape.has_text_frame or visible_fraction < 0.5:
                    overflow["slide"] = "object extends past slide boundary"
            if shape.has_text_frame and shape.text.strip():
                frame = shape.text_frame
                usable_height = max(0, shape.height - frame.margin_top - frame.margin_bottom)
                usable_width = max(1, shape.width - frame.margin_left - frame.margin_right)
                estimated = 0
                for paragraph in frame.paragraphs:
                    if not paragraph.text.strip():
                        continue
                    sizes = [run.font.size.pt for run in paragraph.runs if run.font.size]
                    size = max(sizes, default=16)
                    runs = [run for run in paragraph.runs if run.text.strip()]
                    if runs and all("Pretendard" in (run.font.name or "") for run in runs):
                        # The bundled face is measured, not estimated; its bold is wider than its regular.
                        bold = sum(len(run.text) for run in runs if run.font.bold) * 2 > sum(len(run.text) for run in runs)
                        line_count = max(1, sum(_measured_lines(line, size, bold, usable_width / 12700) for line in paragraph.text.splitlines() if line.strip()))
                    else:
                        approximate_chars = max(1, usable_width / Inches(size / 72 * 0.52))
                        line_count = max(1, sum(max(1, -(-len(line) // approximate_chars)) for line in paragraph.text.splitlines()))
                    estimated += line_count * Inches(size / 72 * 1.15)
                if estimated > usable_height + Inches(0.08):
                    overflow["frame"] = "estimated text height exceeds frame"
            if overflow or off_slide:
                issues[name] = {"overflow": overflow, **({"off_slide": True} if off_slide else {})}
        if issues:
            report[str(slide_number)] = issues
    return report


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("input", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--issues-only", action="store_true")
    args = parser.parse_args()
    args.output.write_text(json.dumps(inspect(args.input), ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
