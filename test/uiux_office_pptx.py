from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'skills/lit-pptx/scripts'))
from craft_extras import analyse  # noqa: E402
from qa_deck import qa  # noqa: E402


class OfficeCraftTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.deck = Presentation()
        self.deck.slide_width = Inches(13.33)
        self.deck.slide_height = Inches(7.5)
        self.slide = self.deck.slides.add_slide(self.deck.slide_layouts[6])

    def check(self, rule):
        path = Path(self.temp.name) / 'case.pptx'
        self.deck.save(path)
        return [finding for finding in analyse(path)['findings'] if finding['rule'] == rule]

    def textbox(self, text, x=1, y=1, w=3, h=1, size=18):
        box = self.slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        box.text = text
        box.text_frame.paragraphs[0].runs[0].font.size = Pt(size)
        return box

    def test_clean_slide(self):
        self.textbox('Example process', w=5, size=32)
        self.textbox('Clear supporting copy.', y=2, w=5)
        self.assertEqual(analyse(self._save())['findings'], [])
        self.assertIn('craft', qa(self._save()))

    def _save(self):
        path = Path(self.temp.name) / 'case.pptx'
        self.deck.save(path)
        return path

    def test_accent_family(self):
        for index, rgb in enumerate([(220, 30, 30), (30, 170, 50), (40, 70, 220)]):
            shape = self.slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1 + index * 2), Inches(2), Inches(1), Inches(1))
            shape.fill.solid()
            shape.fill.fore_color.rgb = RGBColor(*rgb)
        self.assertTrue(self.check('OF-101'))

    def test_numeric_column_alignment(self):
        table = self.slide.shapes.add_table(3, 2, Inches(1), Inches(2), Inches(6), Inches(2)).table
        for row, value in enumerate(['Amount', '120', '340']):
            table.cell(row, 1).text = value
            table.cell(row, 1).text_frame.paragraphs[0].alignment = PP_ALIGN.LEFT
        self.assertTrue(self.check('OF-103'))
        for row in (1, 2):
            table.cell(row, 1).text_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT
        self.assertFalse(self.check('OF-103'))

    def test_body_measure_and_empty_band(self):
        self.textbox('A' * 220, x=.5, y=2, w=12, h=1.5, size=11)
        self.assertTrue(self.check('OF-102'))
        self.assertTrue(self.check('OF-109'))

    def test_concentric_radius(self):
        self.slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1), Inches(2), Inches(4), Inches(3))
        self.slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(2.2), Inches(3.6), Inches(2.6))
        self.assertTrue(self.check('OF-104'))

    def test_explicit_group_ratio(self):
        for x in (1, 6.7):
            panel = self.slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(2), Inches(5.5), Inches(4))
            panel.fill.solid()
            panel.fill.fore_color.rgb = RGBColor(240, 240, 240)
            self.textbox('First line', x=x+.2, y=2.3, w=2, h=.4)
            self.textbox('Second line', x=x+.2, y=3.1, w=2, h=.4)
        self.assertTrue(self.check('OF-105'))

    def test_gradient_glow_and_emoji_bullet(self):
        shape = self.textbox('An example', y=2)
        run = shape.text_frame.paragraphs[0].runs[0]
        run._r.get_or_add_rPr().append(OxmlElement('a:gradFill'))
        shape._element.spPr.get_or_add_effectLst().append(OxmlElement('a:glow'))
        shape.text_frame.paragraphs[0]._p.get_or_add_pPr().append(OxmlElement('a:buChar'))
        shape.text_frame.paragraphs[0]._p.pPr[-1].set('char', '🚀')
        for rule in ('OF-106', 'OF-107', 'OF-108'):
            self.assertTrue(self.check(rule), rule)


    def named(self, shape, name):
        shape.name = name
        return shape

    def test_display_family_and_page_furniture_are_not_hollow_slides(self):
        # A cover named by the compiler is a display slide however many small lines it holds.
        rail = self.slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(4.5), Inches(7.5))
        rail.fill.solid()
        rail.fill.fore_color.rgb = RGBColor(31, 59, 51)
        self.named(rail, 'family@cover-typographic')
        self.named(self.textbox('Quarterly review', x=.4, y=.5, w=3.8, h=.7, size=36), 'title@cover')
        for row in range(4):
            self.textbox(f'Part {row + 1}', x=5, y=.6 + row * .4, w=3, h=.3, size=13)
        self.assertFalse(self.check('OF-109'))
        # The same page named as a content family is judged as content: its empty lower half counts.
        rail.name = 'family@comparison'
        self.assertTrue(self.check('OF-109'))

    def test_rail_touching_two_canvas_edges_is_not_a_card(self):
        rail = self.slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(4.5), Inches(7.5))
        rail.fill.solid()
        rail.fill.fore_color.rgb = RGBColor(31, 59, 51)
        self.named(rail, 'family@comparison')
        self.named(self.textbox('Plan and actual', x=.3, y=.5, w=4, h=1, size=26), 'title@side-rail')
        for row in range(6):
            self.textbox(f'Row {row + 1} with its comparison text', x=4.7, y=.5 + row * 1.0, w=8, h=.8, size=13)
        self.assertFalse(self.check('OF-109'))
        card = self.slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(5), Inches(1), Inches(3), Inches(5))
        card.fill.solid()
        card.fill.fore_color.rgb = RGBColor(240, 240, 240)
        self.textbox('Only the top', x=5.2, y=1.1, w=2.6, h=.4, size=13)
        self.assertTrue(self.check('OF-109'))

    def test_sample_notice_tag_is_not_an_accent(self):
        for index, rgb in enumerate([(31, 59, 51), (203, 211, 218)]):
            shape = self.slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1 + index * 3), Inches(2), Inches(2), Inches(1))
            shape.fill.solid()
            shape.fill.fore_color.rgb = RGBColor(*rgb)
        tag = self.slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(.3), Inches(6.9), Inches(2.6), Inches(.25))
        tag.fill.solid()
        tag.fill.fore_color.rgb = RGBColor(252, 235, 224)
        self.named(tag, 'lit-notice tag')
        self.assertFalse([item for item in self.check('OF-101') if item['severity'] == 'HIGH'])
        tag.name = 'Shape 9'
        self.assertTrue([item for item in self.check('OF-101') if item['severity'] == 'HIGH'])

    def test_outline_around_text_is_a_frame_not_a_leftover_box(self):
        frame = self.slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(.4), Inches(5.3), Inches(12.5), Inches(1.4))
        frame.fill.background()
        frame.line.color.rgb = RGBColor(17, 17, 17)
        self.textbox('Next step: report the plan within two weeks', x=.6, y=5.6, w=6, h=.3, size=13)
        self.textbox('Board review', w=6, size=26)
        kinds = [issue['kind'] for issue in qa(self._save())['composition']['issues']]
        self.assertNotIn('empty_outline', kinds)
        empty = self.slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(7), Inches(2), Inches(4), Inches(2))
        empty.fill.background()
        empty.line.color.rgb = RGBColor(17, 17, 17)
        kinds = [issue['kind'] for issue in qa(self._save())['composition']['issues']]
        self.assertIn('empty_outline', kinds)

    def test_frame_overflow_is_measured_with_the_bundled_face(self):
        from layout_inventory import inspect
        line = 'the same 20 L reactor; only three feed valves are added'
        box = self.textbox(line, x=1, y=2, w=5.08, h=.23, size=13)
        box.text_frame.paragraphs[0].runs[0].font.name = 'Pretendard'
        for side in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'):
            setattr(box.text_frame, side, 0)
        self.assertEqual(inspect(str(self._save())), {})
        box.width = Inches(2.4)
        self.assertIn('frame', inspect(str(self._save()))['1']['shape_1']['overflow'])


    def test_empty_band_follows_the_named_title_wherever_it_stands(self):
        # A bottom title closes the body from below; a table that starts at the top margin is content.
        marker = self.slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, 0, 0)
        self.named(marker, 'family@table-insight')
        table = self.slide.shapes.add_table(6, 3, Inches(.33), Inches(.5), Inches(12.6), Inches(2.7)).table
        table.cell(0, 0).text = 'Item'
        self.textbox('Takeaway one about the table', x=.5, y=3.4, w=6, h=.3, size=13)
        self.textbox('Takeaway two about the table', x=.5, y=4.0, w=6, h=.3, size=13)
        self.textbox('Takeaway three about the table', x=.5, y=4.7, w=6, h=.3, size=13)
        self.named(self.textbox('Revenue by quarter', x=.33, y=5.6, w=9.4, h=1.1, size=26), 'title@bottom-anchor')
        self.assertFalse(self.check('OF-109'))
        for shape in list(self.slide.shapes)[2:5]:
            shape._element.getparent().remove(shape._element)
        self.assertTrue(self.check('OF-109'))

    def test_overlay_plate_on_a_full_page_picture_is_a_display_device(self):
        marker = self.slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.33), Inches(7.5))
        self.named(marker, 'family@image-full')
        plate = self.slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(4.5), Inches(7.9), Inches(2.7))
        plate.fill.solid()
        plate.fill.fore_color.rgb = RGBColor(17, 17, 17)
        self.named(self.textbox('Table set for four', x=.33, y=4.8, w=7.2, h=1, size=26), 'title@overlay')
        self.textbox('Figure 1. The table opened for four', x=.33, y=6.0, w=7.2, h=.2, size=11)
        self.assertFalse(self.check('OF-109'))


class DeckOutputTests(unittest.TestCase):
    """OF-110..OF-114 read the compiled deck as a whole."""

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)

    def flat_deck(self, legacy=False):
        deck = Presentation()
        deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
        deck.core_properties.subject = 'lit-pptx tonality=ledger density=10 grid=compact variance=6'
        for number in range(1, 10):
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            marker = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(.33), Inches(1.0), Inches(2), Inches(.04))
            marker.name = 'family@cover-typographic' if number == 1 else 'family@table-insight'
            x = Inches(60 / 72) if number == 5 else Inches(48 / 72)
            title = slide.shapes.add_textbox(x, Inches(36 / 72), Inches(8), Inches(.5))
            title.text = 'Quarterly revenue' if number != 3 else 'Revenue grew in every quarter'
            title.text_frame.paragraphs[0].runs[0].font.size = Pt(26)
            title.name = 'title@legacy' if legacy else ('title@side-rail' if number == 7 else 'title@top-rule')
            rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, Inches(80 / 72), Inches(8), Inches(2 / 72))
            rule.fill.solid()
            rule.fill.fore_color.rgb = RGBColor(17, 17, 17)
            rule.name = f'Rule {number}'
            table = slide.shapes.add_table(3, 3, Inches(.67), Inches(1.4), Inches(12), Inches(1.0)).table
            table.cell(0, 0).text = 'Item'
        path = Path(self.temp.name) / 'flat.pptx'
        deck.save(path)
        return path

    def test_a_deck_of_one_look_repeated_fails_each_output_check(self):
        findings = analyse(self.flat_deck())['findings']
        high = {item['rule'] for item in findings if item['severity'] == 'HIGH'}
        for rule in ('OF-110', 'OF-111', 'OF-112', 'OF-113', 'OF-114'):
            self.assertIn(rule, high, rule)
        self.assertTrue(any(item['rule'] == 'OF-110' and item['slide'] == 7 and 'named side-rail' in item['value'] for item in findings))
        self.assertTrue(any(item['rule'] == 'OF-113' and item['slide'] == 5 for item in findings))
        self.assertTrue(all(item.get('hint') for item in findings if item['rule'] in ('OF-110', 'OF-111', 'OF-112', 'OF-113', 'OF-114')))

    def test_a_legacy_template_reports_variety_as_advice(self):
        findings = analyse(self.flat_deck(legacy=True))['findings']
        output = [item for item in findings if item['rule'] in ('OF-110', 'OF-111', 'OF-112', 'OF-114')]
        self.assertTrue(output)
        self.assertFalse([item for item in output if item['severity'] == 'HIGH'])

    def test_titles_are_labels_not_sentences(self):
        from deck_output import is_sentence
        claims = ['맞벌이 가구는 빨래에 주말 반나절을 쓴다', '18개월 운영 자금 30억 원을 요청드립니다', '재주문 비율 절반 돌파함',
                  '결제 연동 일정 확정됨', '전년보다 매출이 늘었어요', '계획보다 4%p 뒤처져 있음', 'Revenue grew 12% in Q2',
                  'The staged feed doubles conversion.', 'Retention is the constraint']
        labels = ['분기별 매출 추이', '운영 자금 30억 원 요청 내역 (18개월)', '부가세 포함 단가', '3분기 경영 실적 보고', '왜 지금인가?',
                  'Conversion by feed rate, runs 1-12', 'Method notes', 'Quarterly revenue and margin, 2025-2026']
        self.assertEqual([is_sentence(text) for text in claims], [True] * len(claims))
        self.assertEqual([is_sentence(text) for text in labels], [False] * len(labels))


if __name__ == '__main__':
    unittest.main()
